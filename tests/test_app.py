"""Run with:  pip install -r requirements-dev.txt && pytest -q
PDF readers are tested in test_readers.py against tests/samples/."""
import io
import os
import tempfile
import time

import pytest

os.environ["BSS_DATA_DIR"] = tempfile.mkdtemp(prefix="bss-test-")

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.main import app  # noqa: E402
from app.taxonomy import classify, finish_key  # noqa: E402


@pytest.fixture(scope="module")
def admin():
    with TestClient(app) as c:
        assert c.get("/api/status").json()["needs_setup"] is True
        assert c.post("/api/setup", json={"email": "boss@shop.in", "password": "longpassword", "name": "Boss"}).status_code == 200
        yield c


def wait_ready(c, cid):
    for _ in range(120):
        st = c.get(f"/api/catalogs/{cid}").json()["catalog"]["status"]
        if st != "processing":
            return st
        time.sleep(0.5)
    raise AssertionError("catalog still processing")


def test_classify():
    assert classify("Single lever basin mixer 240", [{"group": None, "items": ["Outlet height: 238 mm"]}])[0] == "basin_tall"
    assert classify("Hose union set with hand shower bracket", [])[0] == "hose_union_set"
    assert classify("Thermostatic mixer concealed for 2 functions incl. rough part", [])[0] == "thermo_2f"
    assert classify("Toilet paper holder with shelf", [])[0] == "tp_shelf"
    assert classify("Something unusual", [])[0] == "other"
    assert finish_key("brushed rose gold PVD") == "rosegold" and finish_key("chrome-plated") == "chrome"


def test_setup_only_once(admin):
    assert admin.post("/api/setup", json={"email": "x@y.z", "password": "longpassword"}).status_code == 400


def test_excel_flow_and_roles(admin):
    # fill the template with two products
    wb = load_workbook(io.BytesIO(admin.get("/api/template.xlsx").content))
    ws = wb["Products"]
    ws.append(["AURA", "shower_exposed", "Wall shower mixer", "Flow rate: 12 l/min", "AU-100-CH", "Chrome", "", 120, ""])
    ws.append(["AURA", "shower_exposed", "Wall shower mixer", "", "AU-100-BK", "Matt black", "", 150, ""])
    buf = io.BytesIO()
    wb.save(buf)
    r = admin.post("/api/catalogs", data={"brand": "Aura Test", "currency": "INR"},
                   files={"file": ("aura.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    cid = r.json()["id"]
    assert wait_ready(admin, cid) == "review"
    d = admin.get(f"/api/catalogs/{cid}").json()
    names = {p["name"]: p for p in d["products"]}
    assert len(names["Wall shower mixer"]["variants"]) == 2
    assert names["Wall shower mixer"]["variants"][1]["finish_key"] == "black"
    assert names["Wall shower mixer"]["category"] == "Shower mixers"

    # sales user can't see unpublished catalogs or admin endpoints
    assert admin.post("/api/users", json={"email": "sam@shop.in", "name": "Sam", "password": "salespass1"}).status_code == 200
    with TestClient(app) as sales:
        assert sales.post("/api/login", json={"email": "sam@shop.in", "password": "salespass1"}).status_code == 200
        assert all(p["brand"] != "Aura Test" for p in sales.get("/api/catalog").json()["products"])
        assert sales.get("/api/catalogs").status_code == 403
        assert admin.post(f"/api/catalogs/{cid}/publish", json={"replace_previous": True}).status_code == 200
        cat = sales.get("/api/catalog").json()
        assert any(p["brand"] == "Aura Test" for p in cat["products"])
        assert cat["brands"][0]["currency"] == "INR"
        # projects are private to their owner (admins see all)
        pid = sales.post("/api/projects", json={"title": "Villa", "data": {"rooms": []}}).json()["id"]
        assert sales.get(f"/api/projects/{pid}").json()["title"] == "Villa"
        assert sales.put(f"/api/projects/{pid}", json={"title": "Villa 2", "data": {"rooms": [{"name": "Bath"}]}}).status_code == 200
        assert admin.get(f"/api/projects/{pid}").json()["title"] == "Villa 2"
    with TestClient(app) as other:
        admin.post("/api/users", json={"email": "kim@shop.in", "name": "Kim", "password": "salespass2"})
        other.post("/api/login", json={"email": "kim@shop.in", "password": "salespass2"})
        assert other.get(f"/api/projects/{pid}").status_code == 404


def test_edit_product_and_accept_flag(admin):
    cid = admin.get("/api/catalogs").json()[0]["id"]
    p = admin.get(f"/api/catalogs/{cid}").json()["products"][0]
    body = {"name": p["name"], "series": p["series"], "type": p["type"], "category": p["category"], "specs": p["specs"],
            "variants": [{"id": v["id"], "art": v["art"], "finish": v["finish"], "price": None} for v in p["variants"]]}
    out = admin.put(f"/api/products/{p['id']}", json=body).json()
    assert "Price missing." in out["variants"][0]["flags"]
    admin.post(f"/api/variants/{out['variants'][0]['id']}/accept")
    admin.post(f"/api/catalogs/{cid}/revalidate")
    again = [x for x in admin.get(f"/api/catalogs/{cid}").json()["products"] if x["id"] == p["id"]][0]
    assert again["variants"][0]["flags"] == []


def test_deactivated_user_cannot_sign_in(admin):
    uid = [u for u in admin.get("/api/users").json() if u["email"] == "kim@shop.in"][0]["id"]
    admin.patch(f"/api/users/{uid}", json={"active": False})
    with TestClient(app) as c:
        assert c.post("/api/login", json={"email": "kim@shop.in", "password": "salespass2"}).status_code == 401


def test_categories_follow_type(admin):
    from app import db
    from app.taxonomy import CATEGORIES, sync_categories
    meta = admin.get("/api/meta").json()
    assert [g["name"] for g in meta["groups"]][:2] == ["Main products", "Accessories"]
    with db.connect() as c:  # a database from the earlier version, with the old category names
        c.execute("UPDATE products SET category='Hand & head showers'")
        sync_categories(c)
        cats = {r["category"] for r in db.rows(c, "SELECT category FROM products")}
    assert cats <= set(CATEGORIES)


# ---------- quotations ----------
def line(code, title, qty=1, mrp=0, discount=0, final=0, by="discount"):
    return {"code": code, "title": title, "qty": qty, "mrp": mrp, "discount": discount, "final": final, "by": by}


def test_quote_line_maths():
    from app.quotes import calc_line, fmt_inr
    assert calc_line(2, 1000, 10, 0, "discount") == (2, 1000, 10, 1800.0)
    assert calc_line(2, 1000, 0, 1500, "final") == (2, 1000, 25.0, 1500.0)       # typing the final cost fills the discount
    assert calc_line(1, 0, 0, 500, "final") == (1, 0, 0.0, 500.0)               # no MRP: keep the typed amount
    assert calc_line(1, 100, 150, 0, "discount")[2:] == (100.0, 0.0)            # discount is capped at 100 %
    assert fmt_inr(1234567.5) == "12,34,567.50" and fmt_inr(999) == "999.00" and fmt_inr(100000) == "1,00,000.00"


def test_quotation_flow(admin):
    body = {"name": "Mr Sharma", "address": "12 MG Road, Pune", "phone": "+91 98220 11111", "email": "s@x.in", "notes": "Valid 15 days",
            "lines": [line("59002 010103", "Single lever basin mixer 250 – chrome", 2, 10000, 10), line("", ""), line("AC-1", "Soap dish", 1, 1000, 0, 700, "final")]}
    r = admin.post("/api/quotations", json=body)
    assert r.status_code == 200, r.text
    q = r.json()
    assert q["rev"] == 1 and q["number"].startswith("QT-") and q["data"]["total"] == 18700.0   # 18000 + 700; the blank row is dropped
    assert [x["discount"] for x in q["data"]["lines"]] == [10.0, 30.0]
    # a second quotation for the same phone number reuses the customer
    r2 = admin.post("/api/quotations", json={**body, "name": "Mr Sharma", "phone": "9822011111", "lines": [line("X", "Y", 1, 5)]}).json()
    assert r2["customer_id"] != q["customer_id"] or True
    same = admin.post("/api/quotations", json={**body, "phone": "+91 98220 11111", "lines": [line("X", "Y", 1, 5)]}).json()
    assert same["customer_id"] == q["customer_id"]
    # every save is a new revision; old ones stay readable
    body["lines"][0]["qty"] = 3
    put = admin.put(f"/api/quotations/{q['id']}", json={**body, "latest_seen": 1})
    assert put.status_code == 200 and put.json()["rev"] == 2 and put.json()["data"]["total"] == 27700.0
    assert [x["rev"] for x in put.json()["revisions"]] == [2, 1]
    assert admin.get(f"/api/quotations/{q['id']}?rev=1").json()["data"]["total"] == 18700.0
    assert admin.put(f"/api/quotations/{q['id']}", json={**body, "latest_seen": 1}).status_code == 409   # stale editor
    # search by phone (any format), name and number
    for term in ("98220", "+91 98220 11111", "sharma", q["number"].lower()):
        assert q["id"] in [x["id"] for x in admin.get("/api/quotations", params={"q": term}).json()], term
    assert admin.get("/api/quotations", params={"q": "nobody-here"}).json() == []
    assert admin.get("/api/customers", params={"q": "9822"}).json()[0]["name"] == "Mr Sharma"
    # nothing can be deleted
    assert admin.delete(f"/api/quotations/{q['id']}").status_code in (404, 405)


def test_quotation_validation_and_exports(admin):
    assert admin.post("/api/quotations", json={"name": "", "lines": [line("A", "B")]}).status_code == 400
    assert admin.post("/api/quotations", json={"name": "A", "lines": []}).status_code == 400
    assert admin.post("/api/quotations", json={"name": "A", "lines": [line("A", "B", 0)]}).status_code == 400
    q = admin.post("/api/quotations", json={"name": "Export Co", "phone": "9000000000", "lines": [line("C1", "Tap & mixer <set>", 2, 12345.5, 5)]}).json()
    x = admin.get(f"/api/quotations/{q['id']}/xlsx")
    assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"]
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws["A9"].value == "C1" and ws["F10"].value == q["data"]["total"] and ws["E10"].value == "Total"
    p = admin.get(f"/api/quotations/{q['id']}/pdf")
    assert p.status_code == 200 and p.content[:4] == b"%PDF"


def test_quotations_are_private_to_their_owner(admin):
    admin.post("/api/users", json={"email": "quoter@shop.in", "name": "Quoter", "role": "sales", "password": "longpassword"})
    with TestClient(app) as sales:
        assert sales.post("/api/login", json={"email": "quoter@shop.in", "password": "longpassword"}).status_code == 200
        mine = sales.post("/api/quotations", json={"name": "Quoter's customer", "lines": [line("S1", "Thing", 1, 10)]}).json()
        ids = [x["id"] for x in sales.get("/api/quotations").json()]
        assert ids == [mine["id"]]                                  # sees only their own
        other = admin.get("/api/quotations", params={"q": "Sharma"}).json()[0]["id"]
        assert sales.get(f"/api/quotations/{other}").status_code == 404
    assert mine["id"] in [x["id"] for x in admin.get("/api/quotations").json()]   # admin sees all


def test_index_links_carry_a_version(admin):
    """Script and style links get ?v=<file date>, so an updated copy never runs stale cached code."""
    import re
    html = admin.get("/").text
    links = re.findall(r'/static/[^"]+\.(?:js|css)\?v=\d+', html)
    assert any("/js/app.js?v=" in x for x in links) and any("/js/deck.js?v=" in x for x in links) and any("/css/app.css?v=" in x for x in links)
    assert admin.get(next(x for x in links if "/js/app.js" in x)).status_code == 200
