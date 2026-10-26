"""Quotations: customers, quotations with saved revisions, Excel and PDF export. Amounts are always INR.

A quotation is never deleted. Every Save adds a revision (data is a JSON snapshot); the quotation row only
keeps the latest revision number and total so lists and search stay fast.

Line maths (the browser mirrors this, the server is the one that counts):
    gross  = qty x MRP
    final  = gross less discount %      (Final cost is the line total, so the table total is the sum of Final cost)
    typing a Final cost instead (by="final") keeps that amount and works the discount % out of it.
"""
import io
import json
import re
from datetime import datetime

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from . import auth, db

router = APIRouter()
CURRENCY = "INR"


# ---------- maths ----------
def money(x):
    return round(float(x) + 1e-9, 2)


def calc_line(qty, mrp, discount, final, by):
    """Return (qty, mrp, discount %, final) made consistent. `by` says which of discount / final the user typed last."""
    qty = max(0.0, float(qty or 0))
    mrp = max(0.0, float(mrp or 0))
    gross = qty * mrp
    if by == "final":
        final = max(0.0, float(final or 0))
        discount = min(100.0, max(0.0, (1 - final / gross) * 100)) if gross > 0 else 0.0
    else:
        discount = min(100.0, max(0.0, float(discount or 0)))
        final = gross * (1 - discount / 100)
    return qty, mrp, round(discount, 2), money(final)


def fmt_inr(n):
    """1234567.5 -> '12,34,567.50' (Indian digit grouping)."""
    s = f"{abs(float(n)):.2f}"
    whole, frac = s.split(".")
    if len(whole) > 3:
        head, tail = whole[:-3], whole[-3:]
        head = re.sub(r"(\d)(?=(\d\d)+$)", r"\1,", head)
        whole = head + "," + tail
    return ("-" if float(n) < 0 else "") + whole + "." + frac


def digits(s):
    return re.sub(r"\D", "", s or "")


# ---------- models ----------
class LineIn(BaseModel):
    code: str = ""
    title: str = ""
    qty: float = 1
    mrp: float = 0
    discount: float = 0
    final: float = 0
    by: str = "discount"


class QuoteIn(BaseModel):
    customer_id: int | None = None
    name: str = ""
    address: str = ""
    phone: str = ""
    email: str = ""
    notes: str = ""
    lines: list[LineIn] = Field(default_factory=list)
    latest_seen: int | None = None  # the newest revision the browser knew about when it loaded


def clean(body: QuoteIn):
    name = body.name.strip()
    if not name:
        raise HTTPException(400, "Enter the customer name.")
    lines = []
    for ln in body.lines:
        code, title = ln.code.strip(), ln.title.strip()
        if not code and not title:
            continue  # blank row
        if ln.qty <= 0:
            raise HTTPException(400, f"Quantity must be more than 0 (row: {title or code}).")
        qty, mrp, disc, final = calc_line(ln.qty, ln.mrp, ln.discount, ln.final, ln.by)
        lines.append({"code": code, "title": title, "qty": qty, "mrp": mrp, "discount": disc, "final": final, "by": "final" if ln.by == "final" else "discount"})
    if not lines:
        raise HTTPException(400, "Add at least one product row.")
    total = money(sum(x["final"] for x in lines))
    customer = {"name": name, "address": body.address.strip(), "phone": body.phone.strip(), "email": body.email.strip()}
    return customer, {"customer": customer, "lines": lines, "notes": body.notes.strip(), "total": total, "currency": CURRENCY}


def now():
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def save_customer(c, customer, customer_id):
    """Update the chosen customer, else reuse one with the same phone number, else add a new one."""
    d = digits(customer["phone"])
    if customer_id and db.row(c, "SELECT id FROM customers WHERE id=?", customer_id):
        cid = customer_id
    else:
        found = db.row(c, "SELECT id FROM customers WHERE phone_digits=?", d) if len(d) >= 6 else None
        cid = found["id"] if found else None
    if cid:
        c.execute("UPDATE customers SET name=?,address=?,phone=?,phone_digits=?,email=? WHERE id=?",
                  (customer["name"], customer["address"], customer["phone"], d, customer["email"], cid))
        return cid
    return c.execute("INSERT INTO customers(name,address,phone,phone_digits,email) VALUES(?,?,?,?,?)",
                     (customer["name"], customer["address"], customer["phone"], d, customer["email"])).lastrowid


def _own(c, qid, u):
    q = db.row(c, "SELECT * FROM quotations WHERE id=?", qid)
    if not q or (q["owner_id"] != u["id"] and u["role"] != "admin"):
        raise HTTPException(404, "Quotation not found.")
    return q


def _revisions(c, qid):
    return db.rows(c, """SELECT r.rev, r.saved_at, u.name saved_by, json_extract(r.data,'$.total') total
                         FROM quotation_revisions r LEFT JOIN users u ON u.id=r.saved_by WHERE r.quotation_id=? ORDER BY r.rev DESC""", qid)


def _detail(c, q, rev=None):
    rev = rev or q["rev"]
    r = db.row(c, "SELECT * FROM quotation_revisions WHERE quotation_id=? AND rev=?", q["id"], rev)
    if not r:
        raise HTTPException(404, "That revision does not exist.")
    cust = db.row(c, "SELECT * FROM customers WHERE id=?", q["customer_id"])
    return {"id": q["id"], "number": q["number"], "customer_id": q["customer_id"], "latest": q["rev"], "rev": rev, "saved_at": r["saved_at"],
            "created_at": q["created_at"], "customer_now": {k: cust[k] for k in ("name", "address", "phone", "email")},
            "data": json.loads(r["data"]), "revisions": _revisions(c, q["id"])}


# ---------- API ----------
@router.get("/api/customers")
def find_customers(q: str = "", u=Depends(auth.current_user)):
    q = q.strip()
    if not q:
        return []
    like, d = f"%{q.lower()}%", digits(q)
    with db.connect() as c:
        sql = "SELECT id,name,address,phone,email FROM customers WHERE lower(name) LIKE ?"
        args = [like]
        if d:
            sql += " OR phone_digits LIKE ?"
            args.append(f"%{d}%")
        return db.rows(c, sql + " ORDER BY name LIMIT 8", *args)


@router.get("/api/quotations")
def list_quotations(q: str = "", customer_id: int | None = None, u=Depends(auth.current_user)):
    where, args = [], []
    if u["role"] != "admin":
        where.append("qt.owner_id=?")
        args.append(u["id"])
    if customer_id:
        where.append("qt.customer_id=?")
        args.append(customer_id)
    q = q.strip()
    if q:
        d = digits(q)
        part = "(lower(cu.name) LIKE ? OR lower(qt.number) LIKE ?" + (" OR cu.phone_digits LIKE ?" if d else "") + ")"
        where.append(part)
        args += [f"%{q.lower()}%", f"%{q.lower()}%"] + ([f"%{d}%"] if d else [])
    sql = """SELECT qt.id, qt.number, qt.rev, qt.total, qt.created_at, qt.updated_at, cu.id customer_id, cu.name, cu.phone, u.name owner
             FROM quotations qt JOIN customers cu ON cu.id=qt.customer_id JOIN users u ON u.id=qt.owner_id"""
    if where:
        sql += " WHERE " + " AND ".join(where)
    with db.connect() as c:
        return db.rows(c, sql + " ORDER BY qt.updated_at DESC, qt.id DESC LIMIT 200", *args)


@router.post("/api/quotations")
def create_quotation(body: QuoteIn, u=Depends(auth.current_user)):
    customer, data = clean(body)
    with db.connect() as c:
        cid = save_customer(c, customer, body.customer_id)
        t = now()
        qid = c.execute("INSERT INTO quotations(customer_id,owner_id,rev,total,created_at,updated_at) VALUES(?,?,1,?,?,?)",
                        (cid, u["id"], data["total"], t, t)).lastrowid
        c.execute("UPDATE quotations SET number=? WHERE id=?", (f"QT-{t[:4]}-{qid:04d}", qid))
        c.execute("INSERT INTO quotation_revisions(quotation_id,rev,data,saved_by,saved_at) VALUES(?,?,?,?,?)", (qid, 1, json.dumps(data), u["id"], t))
        return _detail(c, db.row(c, "SELECT * FROM quotations WHERE id=?", qid))


@router.get("/api/quotations/{qid}")
def get_quotation(qid: int, rev: int | None = None, u=Depends(auth.current_user)):
    with db.connect() as c:
        return _detail(c, _own(c, qid, u), rev)


@router.put("/api/quotations/{qid}")
def save_revision(qid: int, body: QuoteIn, u=Depends(auth.current_user)):
    customer, data = clean(body)
    with db.connect() as c:
        q = _own(c, qid, u)
        if body.latest_seen is not None and body.latest_seen != q["rev"]:
            raise HTTPException(409, f"Revision {q['rev']} was saved by someone else while you were editing. Reopen the quotation to see it.")
        cid = save_customer(c, customer, q["customer_id"])
        rev, t = q["rev"] + 1, now()
        c.execute("UPDATE quotations SET customer_id=?, rev=?, total=?, updated_at=? WHERE id=?", (cid, rev, data["total"], t, qid))
        c.execute("INSERT INTO quotation_revisions(quotation_id,rev,data,saved_by,saved_at) VALUES(?,?,?,?,?)", (qid, rev, json.dumps(data), u["id"], t))
        return _detail(c, db.row(c, "SELECT * FROM quotations WHERE id=?", qid))


# no DELETE route on purpose: quotations and their revisions are kept for good.


# ---------- export ----------
def _file(c, qid, rev, u):
    d = _detail(c, _own(c, qid, u), rev)
    d["dealer"] = db.get_setting(c, "dealer")
    return d


def _stamp(d):
    return f"{d['number']}-rev{d['rev']}"


def build_xlsx(d):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
    cu, lines = d["data"]["customer"], d["data"]["lines"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Quotation"
    bold = Font(bold=True)
    ws["A1"] = d.get("dealer") or "Quotation"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Quotation {d['number']} · revision {d['rev']} · {d['saved_at']}"
    ws["A4"], ws["B4"] = "Customer", cu["name"]
    ws["A5"], ws["B5"] = "Address", cu["address"]
    ws["A6"], ws["B6"] = "Phone", cu["phone"]
    for r in (4, 5, 6):
        ws.cell(r, 1).font = bold
    head = ["Material code", "Product title", "Qty", "MRP (INR)", "Discount %", "Final cost (INR)"]
    for i, h in enumerate(head, 1):
        cell = ws.cell(8, i, h)
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor="E1EEEE")
        cell.alignment = Alignment(horizontal="left" if i < 3 else "right", wrap_text=True)
    inr = '[>=10000000]##\\,##\\,##\\,##0.00;[>=100000]##\\,##\\,##0.00;##,##0.00'
    r = 9
    for ln in lines:
        ws.cell(r, 1, ln["code"])
        ws.cell(r, 2, ln["title"]).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 3, ln["qty"])
        ws.cell(r, 4, ln["mrp"]).number_format = inr
        ws.cell(r, 5, ln["discount"]).number_format = "0.00"
        ws.cell(r, 6, ln["final"]).number_format = inr
        r += 1
    top = Border(top=Side(style="thin"))
    ws.cell(r, 5, "Total").font = bold
    t = ws.cell(r, 6, d["data"]["total"])
    t.font, t.number_format = bold, inr
    for i in range(1, 7):
        ws.cell(r, i).border = top
    if d["data"].get("notes"):
        ws.cell(r + 2, 1, "Notes").font = bold
        ws.cell(r + 2, 2, d["data"]["notes"]).alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEF", (18, 60, 8, 16, 12, 18)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(d):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError:
        raise HTTPException(500, "PDF export needs the 'reportlab' package. Close the app, run start.bat again, or ask your admin to run: pip install -r requirements.txt")
    from xml.sax.saxutils import escape as esc
    cu, lines, total = d["data"]["customer"], d["data"]["lines"], d["data"]["total"]
    ss = getSampleStyleSheet()
    cell = ParagraphStyle("c", parent=ss["Normal"], fontSize=8.5, leading=10.5)
    right = ParagraphStyle("r", parent=cell, alignment=2)
    headst = ParagraphStyle("h", parent=cell, fontName="Helvetica-Bold")
    headr = ParagraphStyle("hr", parent=headst, alignment=2)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=16 * mm,
                            title=f"Quotation {d['number']}", author=d.get("dealer") or "")
    story = [Paragraph(esc(d.get("dealer") or "Quotation"), ParagraphStyle("t", parent=ss["Title"], alignment=0, fontSize=17, spaceAfter=2)),
             Paragraph(f"Quotation <b>{esc(d['number'])}</b> &nbsp;·&nbsp; revision {d['rev']} &nbsp;·&nbsp; {esc(d['saved_at'])}", cell), Spacer(1, 8)]
    who = [f"<b>{esc(cu['name'])}</b>"] + [esc(x) for x in (cu["address"], ("Phone: " + cu["phone"]) if cu["phone"] else "") if x]
    story += [Paragraph("<br/>".join(who), ParagraphStyle("w", parent=cell, fontSize=10, leading=13)), Spacer(1, 10)]
    rows = [[Paragraph(h, headr if i >= 2 else headst) for i, h in enumerate(["Material code", "Product title", "Qty", "MRP (INR)", "Discount", "Final cost (INR)"])]]
    for ln in lines:
        q = ("%g" % ln["qty"])
        rows.append([Paragraph(esc(ln["code"]), cell), Paragraph(esc(ln["title"]), cell), Paragraph(q, right), Paragraph(fmt_inr(ln["mrp"]), right),
                     Paragraph(("%g%%" % ln["discount"]) if ln["discount"] else "-", right), Paragraph(fmt_inr(ln["final"]), right)])
    rows.append(["", "", "", "", Paragraph("Total", headr), Paragraph(fmt_inr(total), headr)])
    tb = Table(rows, colWidths=[28 * mm, 66 * mm, 14 * mm, 24 * mm, 20 * mm, 30 * mm], repeatRows=1)
    tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E1EEEE")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("LINEBELOW", (0, 0), (-1, -2), 0.4, colors.HexColor("#C9D3D5")), ("LINEABOVE", (0, -1), (-1, -1), 0.8, colors.black),
                            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
    story.append(tb)
    if d["data"].get("notes"):
        story += [Spacer(1, 12), Paragraph("<b>Notes</b><br/>" + esc(d["data"]["notes"]).replace("\n", "<br/>"), cell)]
    doc.build(story, onFirstPage=_foot(d), onLaterPages=_foot(d))
    return buf.getvalue()


def _foot(d):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillGray(0.4)
        canvas.drawString(14 * 2.835, 9 * 2.835, f"All amounts in INR · {d['number']} rev {d['rev']}")
        canvas.drawRightString(A4_W - 14 * 2.835, 9 * 2.835, f"Page {doc.page}")
        canvas.restoreState()
    return draw


A4_W = 595.27


def _download(content, name, media):
    return StreamingResponse(io.BytesIO(content), media_type=media, headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.get("/api/quotations/{qid}/xlsx")
def export_xlsx(qid: int, rev: int | None = None, u=Depends(auth.current_user)):
    with db.connect() as c:
        d = _file(c, qid, rev, u)
    return _download(build_xlsx(d), _stamp(d) + ".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@router.get("/api/quotations/{qid}/pdf")
def export_pdf(qid: int, rev: int | None = None, u=Depends(auth.current_user)):
    with db.connect() as c:
        d = _file(c, qid, rev, u)
    return _download(build_pdf(d), _stamp(d) + ".pdf", "application/pdf")
