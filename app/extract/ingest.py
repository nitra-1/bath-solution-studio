"""Turns an uploaded file into products waiting for review.

PDF  -> pdf_parser -> classify -> finish families -> validate -> database (status 'review')
XLSX -> template rows -> same steps
"""
import json
import re
import threading
import traceback
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .. import db
from ..taxonomy import CATEGORIES, TYPE_CAT, TYPES, classify, finish_key
from . import pdf_parser
from .validate import validate

KNOWN_STYLES = {"TEVEO": "Round, soft forms", "TOPIA": "Square, geometric forms"}
TEMPLATE_COLS = ["Series", "Category", "Type", "Product name", "Specs (one per line)", "Article no.", "Finish", "Width mm", "Price", "Photo file name"]


def _log(cat_id, msg):
    with db.connect() as c:
        c.execute("UPDATE catalogs SET log = coalesce(log,'') || ? WHERE id=?", (msg + "\n", cat_id))


def _attributes(specs):
    items = [i for g in specs for i in g["items"]]
    a = {}
    for label, key in [("Outlet height", "outlet_height_mm"), ("Projection", "projection_mm"), ("Height", "height_mm"), ("Length", "length_mm")]:
        for i in items:
            m = re.match(label + r":\s*([\d\- ]+)", i)
            if m:
                a[key] = m.group(1).strip()
                break
    fl = [int(m.group(1)) for i in items for m in [re.search(r"Flow rate[^:]*:\s*(?:limited to )?(\d+)", i)] if m]
    if fl:
        a["flow_l_min"] = max(fl)
    return a


def _clean(t):
    for a, b in [("Withhout", "Without"), ("Intrinsicallysafe", "Intrinsically safe"), ("Flo rate", "Flow rate"), (" u to ", " up to "), ("withg", "with")]:
        t = t.replace(a, b)
    return t


def store_products(con, cat_id, bid, products, profile=None):
    for p in products:
        p["specs"] = [{"group": g["group"], "items": [_clean(i) for i in g["items"]]} for g in p["specs"]]
        if not p.get("type"):
            p["type"], p["name"] = classify(p["title"], p["specs"], p.get("series", ""))
        p["category"] = p.get("category") or TYPE_CAT.get(p["type"], "Other")
        for v in p["variants"]:
            v["finish_key"] = v.get("finish_key") or finish_key(v["finish"])
        p["attributes"] = {**_attributes(p["specs"]), **p.get("attributes", {})}
    validate(products, profile)
    # mark items repeated from another collection
    first = {}
    for p in products:
        arts = tuple(sorted(v["art"] for v in p["variants"]))
        if arts and arts in first and first[arts] != p["series"]:
            p["note"] = f"Same article numbers as the {first[arts]} item; listed in both collections."
        first.setdefault(arts, p["series"])
    for i, p in enumerate(products):
        if p.get("series"):
            db.ensure_series(con, bid, p["series"], KNOWN_STYLES.get(p["series"].upper(), ""))
        pid = con.execute(
            "INSERT INTO products(catalog_id,series,line,category,type,name,specs,attributes,image,source_page,note,sort) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
            (cat_id, p.get("series", ""), p.get("line", ""), p["category"], p["type"], p["name"], json.dumps(p["specs"]),
             json.dumps(p["attributes"]), p.get("image"), p.get("page"), p.get("note"), i)).lastrowid
        for k, v in enumerate(p["variants"]):
            con.execute("INSERT INTO variants(product_id,art,finish,finish_key,width_mm,price,flags,sort) VALUES(?,?,?,?,?,?,?,?)",
                        (pid, v["art"], v["finish"], v["finish_key"], v.get("width_mm"), v.get("price"), json.dumps(v.get("flags", [])), k))


def revalidate(cat_id):
    """Re-run checks after the admin edited a catalog."""
    with db.connect() as c:
        cat = db.row(c, "SELECT * FROM catalogs WHERE id=?", cat_id)
        from .profiles import _load
        mod = _load().get(cat["parser"] or "")
        prof = mod.PROFILE if mod else None
        prods = db.products_for(c, "p.catalog_id=?", cat_id)
        validate(prods, prof)
        for p in prods:
            c.execute("UPDATE products SET attributes=? WHERE id=?", (json.dumps(p["attributes"]), p["id"]))
            for v in p["variants"]:
                c.execute("UPDATE variants SET flags=? WHERE id=?", (json.dumps(v["flags"]), v["id"]))


class NoProducts(Exception):
    pass


def _run(cat_id, path, parser, default_series):
    try:
        with db.connect() as c:
            cat = db.row(c, "SELECT * FROM catalogs WHERE id=?", cat_id)
        bid = cat["brand_id"]
        prefix = f"c{cat_id}"
        if path.suffix.lower() in (".xlsx", ".xlsm"):
            products, prof = read_template(path, default_series), None
            _log(cat_id, f"Read {len(products)} products from the spreadsheet.")
        else:
            _log(cat_id, f"Reading PDF with the '{parser}' reader…")
            products, prof = pdf_parser.parse(str(path), db.IMG_DIR, prefix, parser, default_series, log=lambda m: _log(cat_id, m))
        if not products:
            raise NoProducts("No products found. " + (POOR_HINT.format(reader=parser) if path.suffix.lower() == ".pdf"
                                                    else "Check the spreadsheet uses the template's Products sheet."))
        with db.connect() as c:
            store_products(c, cat_id, bid, products, prof)
            c.execute("UPDATE catalogs SET status='review' WHERE id=?", (cat_id,))
        n_v = sum(len(p["variants"]) for p in products)
        n_f = sum(1 for p in products for v in p["variants"] if v["flags"]) + sum(1 for p in products if p["attributes"].get("issues"))
        _log(cat_id, f"Done: {len(products)} products, {n_v} article numbers, {n_f} items to check.")
        if path.suffix.lower() == ".pdf" and _looks_poor(products):
            _log(cat_id, POOR_HINT.format(reader=parser))
    except NoProducts as e:
        _log(cat_id, "Failed: " + str(e))
        with db.connect() as c:
            c.execute("UPDATE catalogs SET status='failed' WHERE id=?", (cat_id,))
    except Exception as e:  # noqa
        _log(cat_id, "Failed: " + str(e))
        _log(cat_id, traceback.format_exc(limit=3))
        with db.connect() as c:
            c.execute("UPDATE catalogs SET status='failed' WHERE id=?", (cat_id,))


POOR_HINT = ("This PDF didn't read cleanly with the '{reader}' reader. You can fix products by hand here, "
             "or create a reader for this brand: copy the PDF into the project's tests/samples folder, "
             "open Claude Code in the project folder and ask it to \"create a reader for tests/samples/<file>.pdf\". "
             "See 'Adding a new brand layout' in the README.")


def _looks_poor(products):
    """Rough signal that the layout didn't suit the reader."""
    if len(products) < 5:
        return True
    untitled = sum(1 for p in products if p["title"] == "Untitled product")
    unknown = sum(1 for p in products if p.get("type") == "other")
    no_finish = sum(1 for p in products for v in p["variants"] if not v["finish"])
    n_v = max(1, sum(len(p["variants"]) for p in products))
    return untitled / len(products) > 0.2 or unknown / len(products) > 0.4 or no_finish / n_v > 0.3


def start(cat_id, path: Path, parser="generic", default_series=""):
    t = threading.Thread(target=_run, args=(cat_id, path, parser, default_series), daemon=True)
    t.start()
    return t


# ---------- Excel template ----------

def write_template(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(TEMPLATE_COLS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F6E72")
    ws.append(["TEVEO", "Washbasin", "basin_std", "Single lever basin mixer 110", "Outlet height: 104 mm\nProjection: 113 mm\nFlow rate: limited to 6 l/min", "582 04 01 01 00", "chrome-plated", "", 157.50, ""])
    ws.append(["TEVEO", "Washbasin", "basin_std", "Single lever basin mixer 110", "", "582 04 37 01 00", "matt black", "", 186.70, ""])
    widths = [14, 18, 20, 36, 44, 20, 26, 10, 10, 22]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    dv1 = DataValidation(type="list", formula1='"' + ",".join(CATEGORIES) + '"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1="=Types!$A$2:$A$" + str(len(TYPES) + 1), allow_blank=True)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add("B2:B5000")
    dv2.add("C2:C5000")
    t = wb.create_sheet("Types")
    t.append(["Type key", "Meaning", "Category"])
    for k, l, c in TYPES:
        t.append([k, l, c])
    t.column_dimensions["A"].width = 22
    t.column_dimensions["B"].width = 50
    h = wb.create_sheet("How to fill")
    for line in [
        "One row per article number (one finish or size of a product).",
        "Rows with the same Series + Product name become one product with several finishes.",
        "Specs, Category and Type only need to be on the first row of each product.",
        "Type decides how the solution builder uses the product. See the Types sheet; leave blank to let the system guess.",
        "Photos: after upload, add a photo to each product on the review screen.",
        "Prices: numbers only, in the brand's currency (set when uploading).",
    ]:
        h.append([line])
    h.column_dimensions["A"].width = 110
    wb.save(path)


def read_template(path, default_series=""):
    wb = load_workbook(path, data_only=True)
    ws = wb["Products"] if "Products" in wb.sheetnames else wb.active
    hdr = [str(c.value or "").strip().lower() for c in ws[1]]

    def col(name):
        name = name.lower()
        for i, h in enumerate(hdr):
            if h.startswith(name):
                return i
        return None
    ix = {k: col(k) for k in ["series", "category", "type", "product name", "specs", "article", "finish", "width", "price", "photo"]}
    products, by_key = [], {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        g = lambda k: (r[ix[k]] if ix[k] is not None and ix[k] < len(r) else None)
        name = str(g("product name") or "").strip()
        art = str(g("article") or "").strip()
        if not name or not art:
            continue
        series = str(g("series") or default_series or "").strip()
        key = (series, name)
        if key not in by_key:
            spec_txt = str(g("specs") or "")
            typ = str(g("type") or "").strip()
            p = {"series": series, "line": "", "title": name, "name": name, "page": None,
                 "specs": [{"group": None, "items": [s.strip(" •-") for s in spec_txt.splitlines() if s.strip()]}],
                 "variants": [], "image": None}
            if typ in TYPE_CAT:
                p["type"] = typ
                p["category"] = str(g("category") or "").strip() if str(g("category") or "").strip() in CATEGORIES else TYPE_CAT[typ]
            by_key[key] = p
            products.append(p)
        p = by_key[key]
        try:
            price = float(str(g("price")).replace(",", "")) if g("price") not in (None, "") else None
        except ValueError:
            price = None
        w = g("width")
        p["variants"].append({"art": art, "finish": str(g("finish") or "").strip(), "width_mm": int(w) if str(w or "").isdigit() else None, "price": price})
    return products
